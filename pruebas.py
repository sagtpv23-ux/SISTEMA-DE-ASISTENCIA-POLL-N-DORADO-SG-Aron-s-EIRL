"""
Casos de prueba automaticos del prototipo (solo libreria estandar + openpyxl).
Ejecutar:  python pruebas.py
Cada caso compara el resultado esperado con el obtenido y muestra OK / FALLO.
Con --json ruta.json guarda ademas los resultados (usado para el informe).
"""

import json
import os
import sys
import tempfile

from modulo_trabajadores import (
    registrar_empleado,
    ordenar_empleados_por_codigo,
    buscar_binario_recursivo,
)
from modulo_asistencia import (
    marcar_entrada,
    marcar_salida,
    calcular_horas_trabajadas,
    ordenar_asistencias_por_fecha,
    contar_tardanzas_recursivo,
    completar_turnos,
    HORA_TOPE_POR_TURNO,
)
from modulo_planilla import generar_planilla_mensual, total_planilla, contar_jornadas_mes, meses_con_datos
from modulo_persistencia import guardar_empleados, cargar_empleados, guardar_asistencias, cargar_asistencias
from modulo_reportes import exportar_a_excel
from modulo_validaciones import validar_formato_mes

RESULTADOS = []


def caso(codigo, descripcion, esperado, funcion):
    """Ejecuta una prueba; funcion() retorna (obtenido_texto, cumple_bool)."""
    try:
        obtenido, ok = funcion()
    except Exception as error:  # una prueba nunca debe romper a las demas
        obtenido, ok = f"Excepcion: {error}", False
    RESULTADOS.append({"id": codigo, "descripcion": descripcion, "esperado": esperado,
                       "obtenido": obtenido, "ok": ok})
    print(f"[{'OK' if ok else 'FALLO'}] {codigo} {descripcion}")
    print(f"        esperado: {esperado}")
    print(f"        obtenido: {obtenido}")


def base():
    """Estado limpio: 2 trabajadores (uno de cada turno habitual)."""
    empleados, asistencias = [], []
    registrar_empleado(empleados, "E001", "Carlos", "Ramirez", "71234567", "Parrillero", "M", "8.5")
    registrar_empleado(empleados, "E002", "Maria", "Lopez", "75647382", "Mesero", "T", "7.0")
    return empleados, asistencias


def ejecutar():
    # Las pruebas fijan los topes en 08:00 y 14:00 para comprobar la logica, de modo
    # que siguen pasando aunque se cambien las horas reales en modulo_asistencia.py.
    topes_reales = dict(HORA_TOPE_POR_TURNO)
    HORA_TOPE_POR_TURNO.update({"Mañana": "08:00", "Tarde": "14:00"})

    # ---------- Registro y validaciones de trabajadores ----------
    def p01():
        emp, _ = [], []
        ok, msg = registrar_empleado(emp, "E001", "Carlos", "Ramirez", "71234567", "Parrillero", "M", "8.5")
        return msg, ok and emp[0]["turno"] == "Mañana"
    caso("P01", "Registrar trabajador valido (turno M, tarifa 8.5)", "Registro exitoso, turno Mañana", p01)

    def p02():
        emp, _ = base()
        ok, msg = registrar_empleado(emp, "E001", "Otro", "Nombre", "70000000", "Caja", "T", "8")
        return msg, (not ok)
    caso("P02", "Codigo de trabajador duplicado", "Rechazado", p02)

    def p03():
        emp, _ = base()
        ok, msg = registrar_empleado(emp, "X1", "Ana", "Torres", "12345678", "Caja", "T", "8")
        return msg, (not ok)
    caso("P03", "Codigo con formato invalido (X1)", "Rechazado", p03)

    def p04():
        emp, _ = base()
        ok, msg = registrar_empleado(emp, "E003", "Luis", "Perez", "123", "Caja", "T", "8")
        return msg, (not ok)
    caso("P04", "DNI invalido (123)", "Rechazado", p04)

    def p05():
        emp, _ = base()
        ok, msg = registrar_empleado(emp, "E003", "Ana", "Torres", "12345678", "Caja", "Noche", "8")
        return msg, (not ok)
    caso("P05", "Turno inexistente (Noche): la empresa solo tiene 2 turnos", "Rechazado", p05)

    def p06():
        emp, _ = base()
        ok, msg = registrar_empleado(emp, "E003", "Ana", "Torres", "12345678", "Caja", "T", "-5")
        return msg, (not ok)
    caso("P06", "Tarifa por hora negativa (-5)", "Rechazado", p06)

    # ---------- Entrada, tardanza por turno y turno por jornada ----------
    def p07():
        emp, asi = base()
        ok, msg = marcar_entrada(asi, emp, "E001", "2026-09-01", "08:00")
        return f"{msg} tardanza={asi[0]['tardanza']}", ok and asi[0]["tardanza"] == "No"
    caso("P07", "Entrada turno Mañana a las 08:00 (puntual)", "Sin tardanza", p07)

    def p08():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:15")
        return f"tardanza={asi[0]['tardanza']}", asi[0]["tardanza"] == "Si"
    caso("P08", "Entrada turno Mañana a las 08:15", "Tardanza", p08)

    def p09():
        emp, asi = base()
        marcar_entrada(asi, emp, "E002", "2026-09-01", "14:00")
        return f"turno={asi[0]['turno']} tardanza={asi[0]['tardanza']}", asi[0]["tardanza"] == "No"
    caso("P09", "Entrada turno Tarde a las 14:00 (antes marcaba tardanza por error)", "Sin tardanza", p09)

    def p10():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "8:00")
        return f"hora guardada={asi[0]['hora_entrada']} tardanza={asi[0]['tardanza']}", (
            asi[0]["hora_entrada"] == "08:00" and asi[0]["tardanza"] == "No")
    caso("P10", "Hora escrita sin cero inicial (8:00)", "Se normaliza a 08:00, sin tardanza", p10)

    def p11():
        emp, asi = base()
        ok, msg = marcar_entrada(asi, emp, "E001", "2026-09-03", "14:00", "T")
        return f"turno de la jornada={asi[0]['turno']} (habitual: {emp[0]['turno']})", (
            ok and asi[0]["turno"] == "Tarde")
    caso("P11", "Trabajador de Mañana cubre el turno Tarde", "Jornada registrada como Tarde", p11)

    def p12():
        emp, asi = base()
        ok, msg = marcar_entrada(asi, emp, "E999", "2026-09-01", "08:00")
        return msg, (not ok)
    caso("P12", "Entrada de trabajador inexistente (E999)", "Rechazado", p12)

    def p13():
        emp, asi = base()
        ok, msg = marcar_entrada(asi, emp, "E001", "2026-13-45", "08:00")
        return msg, (not ok)
    caso("P13", "Fecha invalida (2026-13-45)", "Rechazado", p13)

    # ---------- Salida y horas ----------
    def p14():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:00")
        ok, msg = marcar_salida(asi, "E001", "2026-09-01", "07:00")
        return msg, (not ok)
    caso("P14", "Hora de salida anterior a la de entrada", "Rechazado", p14)

    def p15():
        emp, asi = base()
        ok, msg = marcar_salida(asi, "E001", "2026-09-01", "17:00")
        return msg, (not ok)
    caso("P15", "Salida sin entrada abierta ese dia", "Rechazado", p15)

    def p16():
        horas = calcular_horas_trabajadas("08:00", "17:30")
        return f"{horas} h", horas == 9.5
    caso("P16", "Calculo de horas trabajadas 08:00 a 17:30", "9.5 h", p16)

    # ---------- Busqueda, ordenamiento y recursividad ----------
    def p17():
        emp, _ = base()
        ordenados = ordenar_empleados_por_codigo(emp)
        hallado = buscar_binario_recursivo(ordenados, "E002")
        no_hallado = buscar_binario_recursivo(ordenados, "E999")
        return f"E002 -> {hallado['nombres']}; E999 -> {no_hallado}", hallado is not None and no_hallado is None
    caso("P17", "Busqueda binaria recursiva por codigo (existente e inexistente)", "E002 hallado, E999 None", p17)

    def p18():
        emp = []
        for cod in ("E003", "E001", "E002"):
            registrar_empleado(emp, cod, "N", "A", "12345678", "Caja", "M", "8")
        orden = [e["codigo"] for e in ordenar_empleados_por_codigo(emp)]
        return str(orden), orden == ["E001", "E002", "E003"]
    caso("P18", "Ordenamiento burbuja de trabajadores por codigo", "['E001','E002','E003']", p18)

    def p19():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-03", "08:00"); marcar_salida(asi, "E001", "2026-09-03", "17:00")
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:00"); marcar_salida(asi, "E001", "2026-09-01", "17:00")
        marcar_entrada(asi, emp, "E002", "2026-09-02", "14:00"); marcar_salida(asi, "E002", "2026-09-02", "22:00")
        fechas = [r["fecha"] for r in ordenar_asistencias_por_fecha(asi)]
        return str(fechas), fechas == ["2026-09-01", "2026-09-02", "2026-09-03"]
    caso("P19", "Ordenamiento por seleccion de asistencias por fecha", "Fechas ascendentes", p19)

    def p20():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:15")   # tardanza
        marcar_entrada(asi, emp, "E002", "2026-09-01", "14:30")   # tardanza
        marcar_entrada(asi, emp, "E001", "2026-09-02", "08:00")   # puntual
        total = contar_tardanzas_recursivo(asi)
        return f"{total} tardanzas", total == 2
    caso("P20", "Conteo recursivo de tardanzas (3 entradas, 2 tarde)", "2", p20)

    # ---------- Planilla mensual ----------
    def p21():
        emp, asi = base()
        for fecha, hi, hs, turno in (("2026-09-01", "08:00", "17:00", ""),
                                     ("2026-09-02", "08:00", "17:00", ""),
                                     ("2026-09-03", "14:00", "22:00", "T")):
            marcar_entrada(asi, emp, "E001", fecha, hi, turno)
            marcar_salida(asi, "E001", fecha, hs)
        plan = generar_planilla_mensual(emp, asi, "2026-09")
        c = plan[0]
        texto = (f"E001: Mañana {c['horas_manana']} h + Tarde {c['horas_tarde']} h = "
                 f"{c['horas_trabajadas']} h x S/ {c['tarifa_hora']} = S/ {c['pago_mensual']}")
        return texto, (c["horas_manana"] == 18.0 and c["horas_tarde"] == 8.0 and c["pago_mensual"] == 221.0)
    caso("P21", "Planilla mensual con horas separadas por turno", "18 h Mañana + 8 h Tarde = 26 h x 8.5 = S/ 221.00", p21)

    def p22():
        return (f"2026-09 -> {validar_formato_mes('2026-09')}; 'septiembre' -> {validar_formato_mes('septiembre')}",
                validar_formato_mes("2026-09") and not validar_formato_mes("septiembre"))
    caso("P22", "Validacion del mes de la planilla (AAAA-MM)", "2026-09 valido, 'septiembre' invalido", p22)

    # ---------- Archivos y excepciones ----------
    def p23():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:00"); marcar_salida(asi, "E001", "2026-09-01", "17:00")
        plan = generar_planilla_mensual(emp, asi, "2026-09")
        with tempfile.TemporaryDirectory() as tmp:
            ruta = exportar_a_excel(emp, asi, os.path.join(tmp, "reporte.xlsx"), planilla=plan)
            from openpyxl import load_workbook
            hojas = load_workbook(ruta).sheetnames
        return f"hojas: {hojas}", hojas == ["Empleados", "Asistencias", "Planilla mensual"]
    caso("P23", "Exportar a Excel (.xlsx)", "3 hojas: Empleados, Asistencias, Planilla mensual", p23)

    def p24():
        emp, asi = base()
        with tempfile.TemporaryDirectory() as tmp:
            ruta = exportar_a_excel(emp, asi, os.path.join(tmp, "no_existe", "reporte.xlsx"))
        return f"retorno={ruta} (el programa no se cerro)", ruta is None
    caso("P24", "Excepcion controlada: no se puede guardar el Excel", "Mensaje de error y retorno None, sin cerrar el programa", p24)

    def p25():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-09-01", "08:00"); marcar_salida(asi, "E001", "2026-09-01", "17:00")
        with tempfile.TemporaryDirectory() as tmp:
            re_, ra_ = os.path.join(tmp, "e.csv"), os.path.join(tmp, "a.csv")
            guardar_empleados(emp, re_); guardar_asistencias(asi, ra_)
            emp2, asi2 = cargar_empleados(re_), cargar_asistencias(ra_)
        return (f"{len(emp2)} trabajadores y {len(asi2)} asistencia(s) recuperados; turno jornada={asi2[0]['turno']}",
                len(emp2) == 2 and len(asi2) == 1 and asi2[0]["turno"] == "Mañana")
    caso("P25", "Persistencia en CSV: guardar y volver a cargar", "Los datos se recuperan completos", p25)

    # ---------- Correcciones: jornadas sin salida, archivos viejos y recursion larga ----------
    def p26():
        n = 5000
        lista = [{"tardanza": "Si" if i % 2 == 0 else "No"} for i in range(n)]
        total = contar_tardanzas_recursivo(lista)
        return f"{n} registros -> {total} tardanzas (sin RecursionError)", total == 2500
    caso("P26", "Conteo recursivo por mitades con 5000 registros", "2500, sin superar el limite de recursion", p26)

    def p27():
        emp, _ = base()
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "asistencias_viejas.csv")
            with open(ruta, "w", encoding="utf-8", newline="") as f:  # version anterior: sin columna turno
                f.write("codigo_empleado,fecha,hora_entrada,hora_salida,horas_trabajadas,tardanza,observaciones\n")
                f.write("E002,2026-10-15,14:00,22:00,8.0,No,\n")
            asi = cargar_asistencias(ruta)
        actualizados = completar_turnos(asi, emp)
        c = generar_planilla_mensual(emp, asi, "2026-10")[1]
        return (f"{actualizados} registro actualizado, turno={asi[0]['turno']}, "
                f"horas Tarde={c['horas_tarde']}, pago=S/ {c['pago_mensual']}",
                actualizados == 1 and asi[0]["turno"] == "Tarde" and c["pago_mensual"] == 56.0)
    caso("P27", "Archivo de una version anterior sin columna turno", "Se asigna el turno habitual y la planilla cuenta las horas", p27)

    def p28():
        emp, asi = base()
        marcar_entrada(asi, emp, "E001", "2026-10-15", "08:00")
        cerradas, abiertas = contar_jornadas_mes(asi, "E001", "2026-10")
        c = generar_planilla_mensual(emp, asi, "2026-10")[0]
        return f"cerradas={cerradas}, abiertas={abiertas}, horas={c['horas_trabajadas']}", (cerradas, abiertas) == (0, 1) and c["horas_trabajadas"] == 0
    caso("P28", "Jornada sin hora de salida: se detecta y la planilla queda en 0", "0 cerradas, 1 abierta, 0 horas", p28)

    def p29():
        emp, asi = base()
        for fecha in ("2026-10-02", "2026-09-30", "2026-10-15"):
            marcar_entrada(asi, emp, "E001", fecha, "08:00")
            marcar_salida(asi, "E001", fecha, "12:00")
        meses = meses_con_datos(asi)
        return str(meses), meses == ["2026-09", "2026-10"]
    caso("P29", "Meses con asistencias registradas (para detectar un mes equivocado)", "['2026-09', '2026-10']", p29)

    HORA_TOPE_POR_TURNO.clear()
    HORA_TOPE_POR_TURNO.update(topes_reales)

    ok_total = sum(1 for r in RESULTADOS if r["ok"])
    print(f"\nRESUMEN: {ok_total}/{len(RESULTADOS)} casos correctos")
    return ok_total == len(RESULTADOS)


if __name__ == "__main__":
    todo_ok = ejecutar()
    if "--json" in sys.argv:
        destino = sys.argv[sys.argv.index("--json") + 1]
        with open(destino, "w", encoding="utf-8") as f:
            json.dump(RESULTADOS, f, ensure_ascii=False, indent=1)
    sys.exit(0 if todo_ok else 1)
